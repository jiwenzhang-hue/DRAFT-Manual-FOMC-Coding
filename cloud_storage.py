"""Server-side Supabase storage and owner-authenticated Excel download."""
from io import BytesIO
import json
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

import streamlit as st


class StorageError(RuntimeError):
    pass


def config():
    try:
        c = dict(st.secrets['supabase'])
        for key in ('url', 'publishable_key', 'secret_key', 'admin_user_id'):
            if not c.get(key) or 'REPLACE' in c[key]:
                raise ValueError(key)
        url = urlparse(c['url'])
        if url.scheme != 'https' or not url.hostname or url.username or url.password:
            raise ValueError('url')
        c['url'] = c['url'].rstrip('/')
        return c
    except Exception:
        raise StorageError('Configure the four supabase values in Streamlit Secrets.') from None


def api(method, path, *, token=None, server_write=False, body=None, prefer=None):
    c = config()
    headers = {'apikey': c['secret_key'] if server_write else c['publishable_key']}
    if token:
        headers['Authorization'] = 'Bearer ' + token
    if prefer:
        headers['Prefer'] = prefer
    data = None
    if body is not None:
        data = json.dumps(body, allow_nan=False).encode('utf-8')
        headers['Content-Type'] = 'application/json'
    request = Request(c['url'] + path, data=data, headers=headers, method=method)
    try:
        with urlopen(request, timeout=30) as response:
            raw = response.read()
            return (json.loads(raw) if raw else None), dict(response.headers)
    except HTTPError as exc:
        if exc.code in (401, 403):
            raise StorageError('Access denied. Check your login or database permissions.') from None
        raise StorageError(f'Database request failed (HTTP {exc.code}). Check the database setup.') from None
    except (URLError, TimeoutError, OSError):
        raise StorageError('Could not reach the database. Your answers remain available; retry submission.') from None


def save_cloud_response(response, submission_id):
    api('POST', '/rest/v1/fomc_responses?on_conflict=id', server_write=True,
        prefer='resolution=ignore-duplicates,return=minimal', body={
            'id': submission_id,
            'start_timestamp': response['start_timestamp'],
            'end_timestamp': response['end_timestamp'],
            'payload': response,
        })


def verify_owner(token):
    user, _ = api('GET', '/auth/v1/user', token=token)
    if not user or user.get('id') != config()['admin_user_id']:
        raise StorageError('This account is not the configured form owner.')
    return user


def login_owner(email, password):
    result, _ = api('POST', '/auth/v1/token?grant_type=password',
                    body={'email': email.strip(), 'password': password})
    token = result.get('access_token') if result else None
    if not token:
        raise StorageError('Sign-in failed.')
    verify_owner(token)
    return token


def fetch_owner_rows(token):
    verify_owner(token)
    rows, offset = [], 0
    while True:
        page, headers = api('GET',
            '/rest/v1/fomc_responses?select=id,created_at,payload'
            f'&order=created_at.asc,id.asc&offset={offset}&limit=500',
            token=token, prefer='count=exact')
        if not page:
            break
        rows.extend({'submission_id': row['id'], 'saved_at_utc': row['created_at'],
                     **row['payload']} for row in page)
        offset += len(page)
        content_range = next((v for k, v in headers.items() if k.lower() == 'content-range'), '')
        total = content_range.rsplit('/', 1)[-1]
        if total.isdigit() and offset >= int(total):
            break
        # If exact counts are unavailable, continue until an empty page.
    return rows


def make_excel(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    columns = list(dict.fromkeys(k for row in rows for k in row))
    if not columns:
        columns = ['submission_id', 'saved_at_utc', 'start_timestamp', 'end_timestamp']
    book = Workbook()
    sheet = book.active
    sheet.title = 'Responses'
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(k, '') for k in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 65
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, name in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 34 if 'timestamp' in name or name == 'saved_at_utc' else 26
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def owner_download_panel():
    with st.sidebar:
        st.header('Owner response downloads')
        if st.session_state.pop('_clear_owner_password', False):
            st.session_state.pop('_owner_password', None)
        token = st.session_state.get('_owner_token')
        if token:
            try:
                verify_owner(token)
            except StorageError:
                st.session_state.pop('_owner_token', None)
                st.session_state.pop('_owner_excel', None)
                st.warning('Please sign in again to access responses.')
                token = None
        if not token:
            with st.form('owner_sign_in'):
                email = st.text_input('Owner email')
                password = st.text_input('Owner password', type='password', key='_owner_password')
                submitted = st.form_submit_button('Sign in')
            if submitted:
                try:
                    st.session_state['_owner_token'] = login_owner(email, password)
                    st.session_state['_clear_owner_password'] = True
                    st.rerun()
                except StorageError as exc:
                    st.error(str(exc))
            return
        st.success('Signed in as the form owner.')
        if st.button('Prepare latest Excel download'):
            try:
                rows = fetch_owner_rows(token)
                st.session_state['_owner_excel'] = make_excel(rows)
                st.session_state['_owner_count'] = len(rows)
            except StorageError as exc:
                st.session_state.pop('_owner_excel', None)
                st.error(str(exc))
        if '_owner_excel' in st.session_state:
            st.caption(f"{st.session_state.get('_owner_count', 0)} submissions")
            st.download_button('Download Excel', st.session_state['_owner_excel'],
                file_name='manual_coding_responses.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        if st.button('Sign out'):
            try:
                api('POST', '/auth/v1/logout', token=token)
            except StorageError:
                pass
            for key in ('_owner_token', '_owner_excel', '_owner_count', '_owner_password'):
                st.session_state.pop(key, None)
            st.rerun()
